from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_qna_workbook(data, output_path="QnA.xlsx"):
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="D9D9D9")
    fill = PatternFill("solid", fgColor="2563EB")

    summary = wb.create_sheet("Summary")
    summary.append(["LinguaQ AI - Q&A Summary",""])
    summary.append(["Language","Q&A Count"])
    for lang in ["English","Hindi","Marathi"]:
        summary.append([lang,len(data.get(lang,[]))])
    summary.merge_cells("A1:B1")
    summary["A1"].font = Font(bold=True,color="FFFFFF",size=16)
    summary["A1"].fill = fill
    for cell in summary[2]:
        cell.font = Font(bold=True,color="FFFFFF")
        cell.fill = fill
    summary.column_dimensions["A"].width=25
    summary.column_dimensions["B"].width=18

    for lang in ["English","Hindi","Marathi"]:
        ws=wb.create_sheet(lang)
        ws.append(["Q.No","Questions","Answers"])
        for i,item in enumerate(data.get(lang,[]),1):
            ws.append([i,item.get("question",""),item.get("answer","")])
        for cell in ws[1]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=fill
            cell.alignment=Alignment(horizontal="center")
            cell.border=Border(bottom=thin)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical="top",wrap_text=True)
                cell.border=Border(bottom=thin)
        ws.column_dimensions["A"].width=10
        ws.column_dimensions["B"].width=55
        ws.column_dimensions["C"].width=85
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
    wb.save(output_path)
    return output_path
